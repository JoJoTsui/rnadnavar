from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "docs/presentation/FASTQ_TO_RESCUED_VCF_OVERVIEW.pptx"
IMG = ROOT / "docs/images/rnadnavar_schemav3.png"
TEMPLATE_CANDIDATES = [
    ROOT / "docs/presentation/agent-template1.pptx",
    Path.home() / "Downloads/agent-template1.pptx",
]


def get_template_path() -> Path | None:
    for candidate in TEMPLATE_CANDIDATES:
        if candidate.exists():
            return candidate
    return None


template_path = get_template_path()
prs = Presentation(str(template_path)) if template_path else Presentation()


def clear_template_slides() -> None:
    """Remove pre-existing template slides so output only contains this deck."""
    slide_ids = list(prs.slides._sldIdLst)
    for slide_id in slide_ids:
        rel_id = slide_id.rId
        prs.part.drop_rel(rel_id)
        prs.slides._sldIdLst.remove(slide_id)


def layout(index: int):
    if len(prs.slide_layouts) > index:
        return prs.slide_layouts[index]
    return prs.slide_layouts[0]


def add_bilingual_notes(slide, en_points: list[str], zh_points: list[str]) -> None:
    """Attach bilingual speaker notes in a consistent EN/CN format."""
    notes = slide.notes_slide.notes_text_frame
    notes.clear()

    p = notes.paragraphs[0]
    p.text = "EN"
    p.font.size = Pt(14)

    for line in en_points:
        ep = notes.add_paragraph()
        ep.text = f"- {line}"
        ep.level = 0
        ep.font.size = Pt(12)

    sep = notes.add_paragraph()
    sep.text = ""

    zh = notes.add_paragraph()
    zh.text = "中文"
    zh.level = 0
    zh.font.size = Pt(14)

    for line in zh_points:
        zp = notes.add_paragraph()
        zp.text = f"- {line}"
        zp.level = 0
        zp.font.size = Pt(12)


def add_title_slide(title: str, subtitle: str, en_notes: list[str], zh_notes: list[str]) -> None:
    slide = prs.slides.add_slide(layout(0))
    slide.shapes.title.text = title
    if len(slide.placeholders) > 1:
        slide.placeholders[1].text = subtitle
    add_bilingual_notes(slide, en_notes, zh_notes)


def add_bullets_slide(
    title: str, bullets: list[str], en_notes: list[str], zh_notes: list[str]
) -> None:
    slide = prs.slides.add_slide(layout(1))
    slide.shapes.title.text = title
    body_idx = 1 if len(slide.placeholders) > 1 else 0
    tf = slide.shapes.placeholders[body_idx].text_frame
    tf.clear()
    for i, line in enumerate(bullets):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = line
        p.level = 0
        p.font.size = Pt(22)
    add_bilingual_notes(slide, en_notes, zh_notes)


def add_workflow_image_slide(en_notes: list[str], zh_notes: list[str]) -> None:
    slide = prs.slides.add_slide(layout(5))
    slide.shapes.title.text = "1) rnadnavar Workflow Introduction"
    if IMG.exists():
        slide.shapes.add_picture(str(IMG), Inches(0.6), Inches(1.2), width=Inches(12.0))
    textbox = slide.shapes.add_textbox(
        Inches(0.7), Inches(6.2), Inches(12.0), Inches(0.8)
    )
    tf = textbox.text_frame
    tf.text = "DNA and RNA paths run in parallel, then cross over in rescue; optional RNA realignment enables second rescue."
    tf.paragraphs[0].font.size = Pt(18)
    add_bilingual_notes(slide, en_notes, zh_notes)


clear_template_slides()

add_title_slide(
    "rnadnavar: From FASTQ to Labeled FILTER",
    "How variant classifications are derived from raw inputs to final VCF FILTER",
    en_notes=[
        "This talk explains the labeling pipeline, not caller internals.",
        "The key goal is traceability: each final FILTER label can be traced to upstream evidence.",
    ],
    zh_notes=[
        "本次汇报重点是标签如何生成，而不是各个 caller 的算法细节。",
        "核心目标是可追溯：最终 FILTER 标签都能追溯到上游证据。",
    ],
)

add_bullets_slide(
    "1) rnadnavar Workflow Introduction",
    [
        "status 0: DNA normal (COO8801DN)",
        "status 1: DNA tumor (COO8801DT)",
        "status 2: RNA tumor (COO8801RT)",
        "Flow: FASTQ -> alignment -> calling -> consensus -> rescue -> final FILTER",
    ],
    en_notes=[
        "Status labels define modality and downstream routing.",
        "DNA and RNA are processed in parallel and converge in rescue.",
    ],
    zh_notes=[
        "status 标签决定样本所属模态及后续流程路由。",
        "DNA 与 RNA 并行处理，并在 rescue 阶段汇合。",
    ],
)

add_workflow_image_slide(
    en_notes=[
        "Use this figure to orient the audience from raw FASTQ to final VCF.",
        "Consensus happens within modality; rescue happens across modalities.",
    ],
    zh_notes=[
        "此图用于展示从 FASTQ 到最终 VCF 的全流程。",
        "consensus 是模态内整合，rescue 是跨模态整合。",
    ],
)

add_bullets_slide(
    "2) How Per-Caller Labels Are Produced",
    [
        "DeepSomatic FILTER -> Somatic/Germline/Reference/Artifact",
        "Mutect2 FILTER -> Somatic/Germline/Reference/Artifact",
        "Strelka FILTER + NT + normal depth -> biological class",
        "These per-caller classes are the inputs to consensus labeling",
    ],
    en_notes=[
        "Each caller output is normalized to the same biological label space.",
        "This standardization enables voting across heterogeneous callers.",
    ],
    zh_notes=[
        "每个 caller 的输出都会先标准化到统一的生物学标签空间。",
        "只有先标准化，后续 consensus 投票才可比较。",
    ],
)

add_bullets_slide(
    "3) Within-Modality Consensus Logics",
    [
        "Group same variant by normalized key within one modality",
        "Apply thresholds: SNV 2, indel 2 (this run)",
        "Clear majority across callers -> consensus label",
        "Top-class tie -> Artifact; below threshold -> NoConsensus",
    ],
    en_notes=[
        "Consensus converts multiple caller labels to one modality-level label.",
        "Tie means disagreement and is intentionally labeled Artifact.",
    ],
    zh_notes=[
        "consensus 将多个 caller 标签汇总为一个模态级标签。",
        "并列代表不一致，因此标记为 Artifact。",
    ],
)

add_bullets_slide(
    "4) Cross-Modality Rescue Logics",
    [
        "Combine DNA and RNA consensus labels for each variant",
        "If labels agree -> keep label; if disagree -> support-aware decision",
        "One-modality consensus can still label variant when supported",
        "Optional second rescue uses realigned RNA consensus",
    ],
    en_notes=[
        "Rescue integrates modality-level evidence, not raw FASTQ directly.",
        "Support counts from individual callers control conflict handling.",
    ],
    zh_notes=[
        "rescue 结合的是模态级证据，而不是直接从 FASTQ 判定。",
        "冲突处理由各模态支持数控制，是 support-aware 逻辑。",
    ],
)

add_bullets_slide(
    "5) Final VCF FILTER Categories",
    [
        "Somatic: final label indicates likely somatic variant",
        "Germline: final label indicates likely inherited variant",
        "Reference: final label indicates non-variant/reference-like evidence",
        "Presentation scope intentionally shows only these three categories",
    ],
    en_notes=[
        "For this audience, final interpretation is focused on three categories.",
        "Other internal classes exist but are not the focus of this deck.",
    ],
    zh_notes=[
        "本次演示只聚焦最终解释最常用的三类 FILTER。",
        "其他内部类别存在，但不作为本次重点。",
    ],
)

add_bullets_slide(
    "6) Traceability: FASTQ to Final FILTER",
    [
        "Input status labels define modality routing",
        "Per-caller mapping creates normalized biological labels",
        "Consensus and rescue propagate and resolve labels",
        "Final VCF FILTER stores the unified final classification",
    ],
    en_notes=[
        "This is the summary chain from raw data to final label.",
        "The message is consistency and traceability across stages.",
    ],
    zh_notes=[
        "该页总结了从原始数据到最终标签的完整链路。",
        "核心信息是各阶段标签定义一致且可追溯。",
    ],
)

prs.save(str(OUT))
print(f"Wrote {OUT}")
